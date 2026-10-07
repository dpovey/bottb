import sys
import qrcode
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT, LOGO, QR_PNG = sys.argv[1], sys.argv[2], sys.argv[3]
URL = 'https://www.battleofthetechbands.com/vote/crowd/sydney-2026'
SHORT = 'battleofthetechbands.com/vote/crowd/sydney-2026'

BG = RGBColor(0x0A, 0x0A, 0x0A)
SURFACE = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xA0, 0xA0, 0xA0)
GOLD = RGBColor(0xF5, 0xA6, 0x23)
FONT = 'Arial'

qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, border=2, box_size=24)
qr.add_data(URL)
qr.make(fit=True)
qr.make_image(fill_color='black', back_color='white').save(QR_PNG)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """runs: list of paragraphs, each a list of (text, size, bold, color)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        for t, size, bold, color in para:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
            r.font.color.rgb = color
    return tb


def rule(s, x, y, w):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD
    bar.line.fill.background()


BANDS = [
    ('Bandlassian', 'Atlassian'),
    ('Canvanauts', 'Canva'),
    ('Amakazaam!', 'Amazon'),
    ('V2 Voyagers', 'V2 AI'),
]

# Slide 1 — crowd voting QR
s = slide()
card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.55), Inches(0.75), Inches(5.2), Inches(6.0))
card.adjustments[0] = 0.04
card.fill.solid()
card.fill.fore_color.rgb = WHITE
card.line.fill.background()
s.shapes.add_picture(QR_PNG, Inches(7.85), Inches(1.0), Inches(4.6), Inches(4.6))
text(s, 7.6, 5.75, 5.1, 0.8, [[(SHORT, 13, False, RGBColor(0x33, 0x33, 0x33))]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

s.shapes.add_picture(LOGO, Inches(0.6), Inches(0.6), height=Inches(1.15))
text(s, 0.6, 2.05, 6.6, 0.5, [[('SYDNEY 2026 · CROWD VOTE', 18, True, GOLD)]])
text(s, 0.6, 2.5, 6.6, 1.3, [[('Vote for your favourite band', 44, True, WHITE)]])
text(s, 0.6, 4.05, 6.6, 1.5, [
    [('1  ', 22, True, GOLD), ('Scan the QR code with your phone camera', 22, False, WHITE)],
    [('2  ', 22, True, GOLD), ('Pick your favourite band and submit', 22, False, WHITE)],
    [('3  ', 22, True, GOLD), ('One vote per person — make it count', 22, False, WHITE)],
])
text(s, 0.6, 5.55, 6.6, 1.3, [
    [('In the running', 16, True, MUTED)],
    [('  ·  '.join(b for b, _ in BANDS), 16, False, MUTED)],
    [('The crowd vote is 20% of the final score.', 16, False, MUTED)],
])

# Slide 2 — run order
s = slide()
s.shapes.add_picture(LOGO, Inches(0.6), Inches(0.5), height=Inches(0.9))
text(s, 0.6, 1.55, 12, 0.9, [[("Tonight's run order", 40, True, WHITE)]])
rule(s, 0.6, 2.45, 2.0)
ROWS = [
    ('7:15', 'ShipReX', 'Rex Software · special guests (not competing)', False),
    ('7:50', 'Bandlassian', 'Atlassian', True),
    ('8:25', 'Canvanauts', 'Canva', True),
    ('9:00', 'Amakazaam!', 'Amazon', True),
    ('9:35', 'V2 Voyagers', 'V2 AI', True),
]
y = 2.8
for t, band, sub, competing in ROWS:
    row = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(y), Inches(12.1), Inches(0.72))
    row.fill.solid()
    row.fill.fore_color.rgb = SURFACE
    row.line.fill.background()
    text(s, 0.85, y, 1.6, 0.72, [[(t, 24, True, GOLD)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.5, y, 4.2, 0.72, [[(band, 26, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 6.6, y, 6.0, 0.72, [[(sub, 18, False, MUTED)]], anchor=MSO_ANCHOR.MIDDLE)
    y += 0.84
text(s, 0.6, 7.0, 12.1, 0.4, [[('Crowd voting: ' + SHORT, 14, False, MUTED)]])

import os
if os.environ.get('ONLY_SLIDE2'):
    ids = prs.slides._sldIdLst
    ids.remove(ids[0])
prs.save(OUT)
print('saved', OUT)
