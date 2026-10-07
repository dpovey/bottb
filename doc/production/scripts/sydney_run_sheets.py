"""Build one single-sheet run sheet per band for BOTTB Sydney 2026.

Replaces Brisbane's input matrix + separate audio and lighting packs. Each band
gets its own file: who plugs in (step 1), own gear (step 2), then the set one
row at a time, songs and the transitions between them (step 3). Instructions
go to bands separately; the festival patch is built from the returned sheets.

Run: uv run --with openpyxl python doc/production/scripts/sydney_run_sheets.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from PIL import Image as PILImage
from PIL import ImageDraw, ImageFont

OUT_DIR = Path(__file__).resolve().parents[1] / "sydney-2026-run-sheets"

EVENT_LINE = "Battle of the Tech Bands Sydney · Thursday 8 October 2026 · Manning Bar, University of Sydney"

# (file-safe name, band, company)
BANDS = [
    ("Bandlassian", "Bandlassian", "Atlassian"),
    ("Canvanauts", "Canvanauts", "Canva"),
    ("Amakazaam", "Amakazaam!", "Amazon"),
    ("V2-Voyagers", "V2 Voyagers", "V2 AI"),
    ("ShipReX", "ShipReX", "Rex Software"),
]

MICS = ["Vox 1", "Vox 2", "Vox 3", "Vox 4", "Vox 5 (drums)"]
FENDER, AC30, AMPEG, KIT = "Fender amp (left)", "AC30 amp (right)", "Ampeg bass amp", "DW kit"
CONNECT = [FENDER, AC30, AMPEG, "Own amp", "DI", "Amp modeller (DI)", *MICS, "Own wireless mic", KIT, "Laptop"]
BACKLINE = (
    'Supplied: DW Collectors kit (22" kick, 10" + 12" racks, 16" floor, 14" snare, double pedal, '
    '3 boom stands, stool) · Zildjian 14" hats, 17" + 19" crashes, 20" ride · Ampeg SVT-VR head + SVT 410 '
    "cab · Fender '65 Super Reverb 4x10 (left) · Vox AC30 2x12 (right) · one single-tier X keyboard stand. "
    "Bring your own snare, cymbals, kick pedal and a second keyboard stand if you need them."
)

PEOPLE_FIRST, PEOPLE_ROWS = 8, 12
PEOPLE_LAST = PEOPLE_FIRST + PEOPLE_ROWS - 1
MAX_SONGS = 7
STEPS = ["Walk-on"]
for n in range(1, MAX_SONGS + 1):
    STEPS.append(f"Song {n}")
    if n < MAX_SONGS:
        STEPS.append(f"Between {n} → {n + 1}")
STEPS.append("Finish")
OTHER_ROW = PEOPLE_LAST + 3
LAYOUT_PNG = OUT_DIR / "stage-layout.png"
RUN_HEADER = OTHER_ROW + 4
RUN_FIRST = RUN_HEADER + 1

WIDTHS = {"A": 16, "B": 28, "C": 20, "D": 22, "E": 26, "F": 26, "G": 30, "H": 44}

INK = "1F2937"
BRAND = "7C3AED"
PALE = "F3F0FF"
BETWEEN_FILL = "FFF7E6"
GREY = "6B7280"
thin = Side(style="thin", color="D1D5DB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")


def section(ws, row, text, hint=""):
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, size=13, color="FFFFFF")
    for col in range(1, 9):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=BRAND)
    if hint:
        ws.cell(row=row, column=3, value=hint).font = Font(italic=True, size=10, color="FFFFFF")
    ws.row_dimensions[row].height = 22


def header_row(ws, row, labels):
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = Font(bold=True, size=10, color=INK)
        c.fill = PatternFill("solid", fgColor=PALE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    ws.row_dimensions[row].height = 45


def box(ws, row, cols=8, fill=None, height=None):
    for col in range(1, cols + 1):
        c = ws.cell(row=row, column=col)
        c.border = BOX
        c.alignment = WRAP
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[row].height = height


def dropdown(ws, options, ref, strict=True):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    # Loose lists still offer the dropdown but let a band type their own answer.
    dv.showErrorMessage = strict
    ws.add_data_validation(dv)
    dv.add(ref)


def stage_layout_png(path):
    """Draw the fixed stage layout, as the audience sees it (Brisbane band-template style)."""
    font = "/System/Library/Fonts/Supplemental/Arial{}.ttf"
    bold = ImageFont.truetype(font.format(" Bold"), 34)
    small = ImageFont.truetype(font.format(""), 24)
    italic = ImageFont.truetype(font.format(" Italic"), 24)
    im = PILImage.new("RGB", (1400, 640), "white")
    d = ImageDraw.Draw(im)

    def block(xy, fill, title, sub=""):
        d.rounded_rectangle(xy, radius=14, fill=fill, outline="#9CA3AF", width=3)
        cx, cy = (xy[0] + xy[2]) / 2, (xy[1] + xy[3]) / 2
        d.text((cx, cy - (14 if sub else 0)), title, font=bold, fill="#1F2937", anchor="mm")
        if sub:
            d.text((cx, cy + 24), sub, font=small, fill="#374151", anchor="mm")

    d.text((700, 32), "UPSTAGE (back wall)", font=italic, fill="#6B7280", anchor="mm")
    block((60, 70, 340, 250), "#FDF1E3", "Fender", "Super Reverb 4x10")
    block((380, 70, 880, 280), "#E8F0FA", "DW drum kit", "double pedal · drummer on Vox 5")
    # Bass sits against the kit so bass and drums can lock in; guitar amps go outside.
    block((920, 70, 1080, 250), "#FDF5D6", "Ampeg", "SVT + 4x10")
    block((1100, 70, 1360, 250), "#FDF1E3", "Vox AC30", "2x12")
    d.text(
        (700, 335), "Keys, special instruments and own amps placed on the day", font=italic, fill="#6B7280", anchor="mm"
    )
    for i, x in enumerate((90, 430, 770, 1110), start=1):
        block((x, 380, x + 210, 480), "#F7E8F2", f"Vox {i}")
    d.text((700, 545), "DOWNSTAGE", font=italic, fill="#6B7280", anchor="mm")
    d.line((60, 585, 1360, 585), fill="#9CA3AF", width=3)
    d.text(
        (700, 612), "AUDIENCE  (left → right as the audience sees the stage)", font=italic, fill="#6B7280", anchor="mm"
    )
    im.save(path)


def run_sheet(band, company, people=(), run=None, other="", anything=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "Run sheet"
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w

    ws["A1"] = f"{band} — run sheet"
    ws["A1"].font = Font(bold=True, size=16, color=INK)
    ws["A2"] = EVENT_LINE
    ws["A2"].font = Font(size=10, color=GREY)

    for i, (k1, v1, k2) in enumerate(
        [("Band / company", f"{band} ({company})", "Set time"), ("Contact on the day", "", "Phone")], start=3
    ):
        for col, val, bold in ((1, k1, True), (2, v1, False), (5, k2, True)):
            ws.cell(row=i, column=col, value=val).font = Font(bold=bold, size=10, color=INK)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
        ws.merge_cells(start_row=i, start_column=6, end_row=i, end_column=8)

    section(ws, PEOPLE_FIRST - 2, "1 — Who's in the band", "One row per person, or per instrument if they swap.")
    header_row(
        ws,
        PEOPLE_FIRST - 1,
        [
            "Name",
            "Plays",
            "Plugs in via\n(Fender, AC30, Ampeg, own amp, DI, Vox 1–5)",
            "Also sings on\n(Vox 1–4 L→R, Vox 5 = drums)",
            "Notes (amp modeller? wireless? in-ears?)",
        ],
    )
    # Notes span E:F. Excel won't auto-fit merged cells, so size rows from their text.
    ws.cell(row=PEOPLE_FIRST - 1, column=6).border = BOX
    ws.merge_cells(start_row=PEOPLE_FIRST - 1, start_column=5, end_row=PEOPLE_FIRST - 1, end_column=6)
    for i, r in enumerate(range(PEOPLE_FIRST, PEOPLE_LAST + 1)):
        person = people[i] if i < len(people) else ()
        for col, val in enumerate(person, start=1):
            ws.cell(row=r, column=col, value=val)
        lines = max([1] + [-(-len(str(v)) // w) for v, w in zip(person, (16, 28, 20, 22, 56))])
        box(ws, r, cols=6, height=max(26, 13 * lines + 4))
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    # Stage layout sits beside the band table, under its own label in the section bar.
    ws.cell(row=PEOPLE_FIRST - 2, column=7, value="Stage layout — as the audience sees it").font = Font(
        bold=True, size=13, color="FFFFFF"
    )
    img = Image(str(LAYOUT_PNG))
    img.width, img.height = 530, 242
    ws.add_image(img, f"G{PEOPLE_FIRST - 1}")
    dropdown(ws, CONNECT, f"C{PEOPLE_FIRST}:C{PEOPLE_LAST}", strict=False)
    dropdown(ws, MICS, f"D{PEOPLE_FIRST}:D{PEOPLE_LAST}")

    section(
        ws,
        OTHER_ROW - 1,
        "2 — Own gear",
        "Own amps, mics, wireless, in-ears, extra inputs. Using an amp modeller? Please level your patches to a similar volume.",
    )
    ws.merge_cells(start_row=OTHER_ROW, start_column=1, end_row=OTHER_ROW, end_column=8)
    c = ws.cell(row=OTHER_ROW, column=1, value=BACKLINE)
    c.font = Font(italic=True, size=10, color=GREY)
    c.alignment = WRAP
    ws.row_dimensions[OTHER_ROW].height = 28
    ws.merge_cells(start_row=OTHER_ROW + 1, start_column=1, end_row=OTHER_ROW + 1, end_column=8)
    box(ws, OTHER_ROW + 1, height=55)
    ws.cell(row=OTHER_ROW + 1, column=1, value=other)

    section(ws, RUN_HEADER - 1, "3 — Your set, in order", "Every song, and what happens between songs.")
    header_row(
        ws,
        RUN_HEADER,
        [
            "Step",
            "Song — artist, or what happens",
            "Length",
            "Who's playing\n('Everyone' or names)",
            "Who sings lead · who moves where",
            "Screen · video · playback audio",
            "Lighting cues or suggestions\n(keep it simple, e.g. blackout to reds and greens)",
            "Other notes",
        ],
    )
    for i, step in enumerate(STEPS):
        r = RUN_FIRST + i
        between = step.startswith("Between")
        box(ws, r, fill=BETWEEN_FILL if between else None, height=30 if between else 40)
        ws.cell(row=r, column=1, value=step).font = Font(bold=not between, size=10, color=GREY if between else INK)
        for col, val in enumerate((run or {}).get(step, ()), start=2):
            ws.cell(row=r, column=col, value=val)

    last = RUN_FIRST + len(STEPS) - 1
    section(ws, last + 2, "4 — Anything else we need to know?", "Requests, surprises, questions for us.")
    ws.merge_cells(start_row=last + 3, start_column=1, end_row=last + 3, end_column=8)
    box(ws, last + 3, height=110)
    ws.cell(row=last + 3, column=1, value=anything)

    ws.freeze_panes = "A3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"
    # Page 1 is the people and gear, page 2 the set and anything else.
    ws.row_breaks.append(Break(id=RUN_HEADER - 2))
    return wb


# ShipReX's Brisbane line-up (Aug 2026 audio pack) as a head start; they confirm or change it.
SHIPREX_PEOPLE = [
    ("Scott", "Vocals", "Vox 3", "", "Lead + backing. From Brisbane — please check"),
    ("Hamish", "Vocals", "Vox 4", "", "Lead + backing. From Brisbane — please check"),
    ("Mindy", "Vocals", "Vox 1", "", "Backing. From Brisbane — please check"),
    ("Elle", "Vocals", "Vox 1", "", "Backing. Shared Vox 1 with Mindy in Brisbane"),
    ("Lil", "Vocals", "Vox 2", "", "Backing. From Brisbane — please check"),
    ("Scotty", "Guitar", FENDER, "", "Amp was mic'd in Brisbane — please check"),
    ("Thiago", "Guitar", AC30, "", "Amp was mic'd in Brisbane — please check"),
    ("Kirra", "Bass", "DI", "", "Recorded as a DI in Brisbane — please check"),
    ("Simon", "Drums", KIT, "Vox 5 (drums)", "From Brisbane — please check"),
    ("", "Erhu", "Mic on instrument", "", "Who plays? Clip-on or wireless?"),
    ("", "Keys", "DI", "", "Who plays?"),
]


# ShipReX's Brisbane set with lighting as they sent it, plus Mustang Sally (new for Sydney).
# Columns: song, length, who's playing, lead / moves, screen & playback, lighting, other notes.
BRISBANE = "As Brisbane — please check"
SHIPREX_RUN = {
    "Song 1": (
        "Uprising — Muse (opens with 30 s of Careless Whisper on erhu)",
        "6:00",
        "Everyone except Scotty (guitar)",
        "",
        "",
        (
            "Dark, spotlight on the erhu to begin. Blackout on the last erhu note, then hit the full band in "
            "deep red on the Uprising riff. Moody red through the verses with a slow pulse on the stomp beat, "
            "brighter white hits on the choruses"
        ),
        BRISBANE,
    ),
    "Song 2": (
        "Espresso — Sabrina Carpenter",
        "3:00",
        "Everyone except Scotty (guitar)",
        "",
        "",
        (
            "Bright and poppy: full warm wash, yellows, oranges and hot pink. No dark moments. Colour sweep or "
            "mirror ball on each chorus; keep the crowd lit"
        ),
        BRISBANE,
    ),
    "Song 3": (
        "Dumb Things — Paul Kelly",
        "3:00",
        "Everyone (Scotty joins on guitar)",
        "",
        "",
        (
            "Warm white pub-rock wash, even and front-lit so faces are visible. Lift the wash and bring up the "
            "crowd for the singalong choruses"
        ),
        BRISBANE,
    ),
    "Song 4": (
        "Mustang Sally — Wilson Pickett",
        "",
        "",
        "",
        "",
        "",
        "New song for Sydney — please fill in this row",
    ),
    "Song 5": (
        "Everlong — Foo Fighters",
        "4:30",
        "Everyone",
        "",
        "",
        (
            "Cool blues and purples on the verses, slam to bright white on every chorus. Drop to a single colour "
            "for the quiet bridge, then full blast for the last chorus"
        ),
        BRISBANE,
    ),
    "Song 6": (
        "Covered in Chrome — Violent Soho",
        "3:45",
        "Everyone",
        "",
        "",
        (
            "Closer, go big: hot white and red, strobes on the chant, everything on for the outro. Hard blackout "
            "on the final hit"
        ),
        BRISBANE,
    ),
}


# Bandlassian's set order (sent 28 Sep 2026); everything else is theirs to fill in.
BANDLASSIAN_RUN = {
    "Song 1": ("good 4 u — Olivia Rodrigo (rock/metal version)",),
    "Song 2": ("Still into You — Paramore",),
    "Song 3": ("Life Is a Highway — Rascal Flatts",),
    "Song 4": ("I'm a Believer — Smash Mouth",),
    "Song 5": ("Pokémon Theme — Jason Paige",),
}

# Amakazaam!'s set and audio notes (sent 28 Sep 2026). The Alexa bits are laptop-audio transitions.
AMAKAZAAM_PEOPLE = [
    (
        "Keira",
        "Lead vocals; drums on Are You Gonna Go My Way",
        "Own wireless mic",
        "Vox 5 (drums)",
        (
            "Stands at Vox 2 (centre). Own wireless mic; its receiver has an XLR out for the desk. Own wireless "
            "in-ears (her transmitter, fed from the desk): her lead vocal + the drum-kit vocal mic when she sings there"
        ),
    ),
    ("Tyler", "", "", "", "Sings on Under Pressure and Killing in the Name. Plays? Which mic?"),
    ("Aaron", "Keys", "", "", "Sings from the keys on Killing in the Name. Which mic?"),
    ("Andrew", "Drums? (please check)", "", "Vox 5 (drums)", "Sings on Are You Gonna Go My Way"),
    ("Laptop", "Alexa audio and video", "Laptop", "", "Run from side stage"),
]
AMAKAZAAM_OTHER = "Keira brings her own wireless mic (receiver with XLR out) and wireless in-ears (own transmitter)."
AMAKAZAAM_ANYTHING = (
    "Request: a boomless (straight) mic stand for the centre mic, usually Vox 2. We can swap stands in the changeover."
)
ALEXA = "Laptop audio + video, run from side stage"
AMAKAZAAM_RUN = {
    "Walk-on": ("Alexa's intro", "", "", "", ALEXA, "", ""),
    "Song 1": ("Valerie — Amy Winehouse", "", "", "Keira", "", "", ""),
    "Song 2": ("Superstition — Stevie Wonder", "", "", "Keira", "", "", ""),
    "Song 3": ("Under Pressure — Queen & David Bowie", "", "", "Keira and Tyler", "", "", ""),
    "Between 3 → 4": ("Alexa's reminder", "", "", "", ALEXA, "", ""),
    "Song 4": (
        "Are You Gonna Go My Way — Lenny Kravitz",
        "",
        "Keira on drums",
        "Drum-kit mic: Andrew, then Keira",
        "",
        "",
        "",
    ),
    "Song 5": ("You Oughta Know — Alanis Morissette", "", "", "Keira", "", "", ""),
    "Between 5 → 6": ("Alexa dialogue", "", "", "", ALEXA, "", ""),
    "Song 6": (
        "Killing in the Name — Rage Against the Machine",
        "",
        "",
        "Keira, Tyler, and Aaron at the keys",
        "",
        "",
        "",
    ),
    "Finish": ("Alexa's outro", "", "", "", ALEXA, "", ""),
}

# Canvanauts' setlist (sent 28 Sep 2026); versions confirmed by Dean.
CANVANAUTS_RUN = {
    "Song 1": ("The Joker and the Thief — Wolfmother",),
    "Song 2": ("Thnks fr th Mmrs — Fall Out Boy",),
    "Song 3": ("So Easy (To Fall in Love) — Olivia Dean",),
    "Song 4": ("Beggin' — Måneskin",),
    "Song 5": ("Smooth Criminal — Alien Ant Farm",),
    "Song 6": ("Welcome to the Black Parade — My Chemical Romance",),
}

# V2 Voyagers' setlist (sent 28 Sep 2026).
V2_VOYAGERS_RUN = {
    "Song 1": ("Venus — Bananarama",),
    "Song 2": ("Cosmic Girl — Jamiroquai",),
    "Song 3": ("Drops of Jupiter — Train",),
    "Song 4": ("All Star — Smash Mouth",),
    "Song 5": ("Medley: Starships — Nicki Minaj / UFO — Sneaky Sound System",),
}

# slug -> (people, run, other, anything) prefills
PREFILL = {
    "ShipReX": (SHIPREX_PEOPLE, SHIPREX_RUN, "", ""),
    "Bandlassian": ((), BANDLASSIAN_RUN, "", ""),
    "Canvanauts": ((), CANVANAUTS_RUN, "", ""),
    "V2-Voyagers": ((), V2_VOYAGERS_RUN, "", ""),
    "Amakazaam": (AMAKAZAAM_PEOPLE, AMAKAZAAM_RUN, AMAKAZAAM_OTHER, AMAKAZAAM_ANYTHING),
}


def main():
    OUT_DIR.mkdir(exist_ok=True)
    stage_layout_png(LAYOUT_PNG)
    for slug, band, company in BANDS:
        out = OUT_DIR / f"BOTTB-Sydney-2026-run-sheet-{slug}.xlsx"
        people, run, other, anything = PREFILL.get(slug, ((), None, "", ""))
        run_sheet(band, company, people=people, run=run, other=other, anything=anything).save(out)
        print(out)


if __name__ == "__main__":
    main()
