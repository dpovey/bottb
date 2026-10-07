"""Build the per-band stage plots, patches and sets for BOTTB Sydney 2026.

One PowerPoint deck that converts cleanly to Google Slides, so the venue techs
and the bands can drag the stage-plot shapes around: schedule, festival patch,
then a stage plot + patch page and a set page for each band in show order.

Band data is transcribed from the run sheets the bands filled in on Drive
(`bottb:Events 2026/Sydney/Run Sheets`, as of 6 Oct 2026). Anything a band
has not told us is marked TBC rather than guessed.

Writes one deck per band (bands don't see each other's sets), a venue
overview (schedule, festival patch, changeovers, links to the band decks on
Drive), and a collated deck of everything for once the bands have confirmed.

Run: uv run --with python-pptx python doc/production/scripts/sydney_stage_plots.py
"""

from math import ceil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

OUT_DIR = Path(__file__).resolve().parents[1] / "sydney-2026-stage-plots"

EVENT = "Battle of the Tech Bands Sydney · Thursday 8 October 2026 · Manning Bar, University of Sydney"
PAGE_W, PAGE_H, MARGIN = 11.69, 8.27, 0.35  # A4 landscape, inches
# Manning Bar stage and drum riser, metres (USU production specs V14, Jul 2024).
STAGE_M, RISER_M = (9.0, 4.6), (2.4, 1.8)
FONT = "Arial"

INK, GREY, RULE = "1F2937", "6B7280", "D1D5DB"
BRAND, PALE, BETWEEN = "7C3AED", "F3F0FF", "FFF7E6"

# kind -> (fill, outline, width m, depth m, label placement). Footprints are real sizes; DI boxes and mic
# stands are drawn a little larger than life so they stay visible.
STYLE = {
    "riser": ("E8F0FA", "9CA3AF", *RISER_M, "inside"),
    "fender": ("FDF1E3", "9CA3AF", 0.63, 0.25, "below"),  # '65 Super Reverb
    "ac30": ("FDF1E3", "9CA3AF", 0.70, 0.27, "below"),  # AC30 C2X
    "ampeg": ("FDF5D6", "9CA3AF", 0.61, 0.45, "below"),  # SVT 410 cab, SVT-VR head on top
    "person": ("FFFFFF", BRAND, 0.5, 0.5, "inside"),
    "mic": ("F7E8F2", "C084B5", 0.3, 0.3, "below"),
    "di": ("FEF3C7", "D97706", 0.22, 0.14, "right"),
    "keys": ("DCFCE7", "16A34A", 1.45, 0.43, "inside"),  # Korg Kronos 2, 88 keys
    "keys_own": ("DCFCE7", "16A34A", 1.2, 0.4, "inside"),
    "laptop": ("E0F2FE", "0284C7", 0.6, 0.45, "below"),
    "iem": ("FFEDD5", "EA580C", 0.5, 0.4, "below"),
}
LEGEND = [
    ("fender", "Backline"),
    ("person", "Player"),
    ("mic", "Vocal mic"),
    ("di", "DI box"),
    ("keys", "Keys"),
    ("laptop", "Laptop / screen"),
    ("iem", "In-ears / wireless"),
]
# Fixed backline, (x, y) centres in metres from house left and from the back wall. Dean's layout
# (6 Oct 2026): riser 0.4 m off the back wall, amps either side of its front half.
RISER_Y = 1.3
BACKLINE = {"fender": (2.45, 1.85), "ampeg": (6.25, 1.75), "ac30": (7.3, 1.8)}
MIC_X, MIC_Y, SINGER_Y = (1.3, 3.4, 5.6, 7.7), 4.4, 3.85

# Festival patch: (channel, source, note). The same for every band.
FESTIVAL = [
    (1, "Kick in", 'House DW kit: 22" kick, 10" + 12" racks, 16" floor, 14" snare'),
    (2, "Kick out", ""),
    (3, "Snare top", ""),
    (4, "Snare bottom", ""),
    (5, "Hi-hat", ""),
    (6, "Rack tom 1", ""),
    (7, "Rack tom 2", ""),
    (8, "Floor tom", ""),
    (9, "Overhead L", ""),
    (10, "Overhead R", ""),
    (11, "Bass DI", "Ampeg SVT-VR DI out, or the player's own DI / modeller XLR"),
    (12, "Bass cab mic", "Ampeg SVT 410 cab"),
    (13, "Gtr 1 (Fender)", "Mic on the Fender '65 Super Reverb, house left"),
    (14, "Gtr 1 DI", "Amp modeller"),
    (15, "Gtr 2 (Vox AC30)", "Mic on the Vox AC30, house right, next to the bass rig"),
    (16, "Gtr 2 DI", "Amp modeller"),
    (17, "Acoustic guitar DI", "V2 Voyagers' acoustic; Canvanauts' third amp modeller"),
    (18, "Instrument DI", "Erhu pickup"),
    (19, "Keys L", "DI"),
    (20, "Keys R", "DI"),
    (21, "Vox 1 (house left)", "Vocal mics numbered house left → house right"),
    (22, "Vox 2", ""),
    (23, "Vox 3", ""),
    (24, "Vox 4 (house right)", ""),
    (25, "Vox 5 (drums)", ""),
    (26, "Wireless vocal", "Amakazaam!'s own receiver, XLR out"),
    (27, "Playback L", "Laptop audio (V2 Voyagers: mono backing track, L only)"),
    (28, "Playback R", ""),
    (29, "Spare", ""),
    (30, "Spare", ""),
    (31, "Crowd L", "Recording only; keep out of FOH"),
    (32, "Crowd R", "Recording only; keep out of FOH"),
]
SOURCE = {ch: src for ch, src, _ in FESTIVAL}

SCHEDULE = [
    ("3:30 pm", "Bump-in", ""),
    ("5:00 pm", "Soundcheck: V2 Voyagers", "10 minutes each, in the reverse of show order"),
    ("5:15 pm", "Soundcheck: Amakazaam!", ""),
    ("5:30 pm", "Soundcheck: Canvanauts", ""),
    ("5:45 pm", "Soundcheck: Bandlassian", ""),
    ("6:00 pm", "Soundcheck: ShipReX", "Stage stays set for ShipReX, who open"),
    ("6:30 pm", "Doors", ""),
    ("7:00–7:15", "Intros", "Host uses one of the band vocal mics"),
    ("7:15–7:40", "ShipReX (Rex Software)", "Special guests"),
    ("7:40–7:50", "Changeover", ""),
    ("7:50–8:15", "Bandlassian (Atlassian)", ""),
    ("8:15–8:25", "Changeover", ""),
    ("8:25–8:50", "Canvanauts (Canva)", ""),
    ("8:50–9:00", "Changeover", ""),
    ("9:00–9:25", "Amakazaam! (Amazon)", ""),
    ("9:25–9:35", "Changeover", ""),
    ("9:35–10:00", "V2 Voyagers (V2 AI)", ""),
]

HOW_TO_READ = [
    (
        "Left and right",
        "Always as the audience sees the stage (house left / house right). Vocal mics Vox 1–4 run house left "
        "to house right along the front; Vox 5 is at the kit.",
    ),
    (
        "One patch all night",
        "Page 2 is the festival patch. Each band's deck lists only the channels that are live for that set; "
        "everything else can stay muted.",
    ),
    (
        "Stage plots are drafts",
        "Built from each band's run sheet, to scale: the stage is 9 m × 4.6 m, the drum riser 2.4 m × 1.8 m, "
        "and amps, keys and players are their real size. Every item is its own shape, so drag it to where it "
        "should go. "
        "TBC marks something the band hasn't told us yet. Grey dashed items aren't used in that set.",
    ),
    (
        "Shared backline",
        "DW Collectors kit with double pedal · Ampeg SVT-VR + SVT 410 · Fender '65 Super Reverb (house left) · "
        "Vox AC30 (house right) · Korg Kronos 2 (Bandlassian and Canvanauts) · one single-tier X keyboard stand.",
    ),
]

FOLD = "fold down"
DRIVE = "https://docs.google.com/presentation/d/{}/edit"
# Band decks on Drive (Events 2026/Sydney Stage Plots (private)), linked from the venue overview.
DECKS = {
    "ShipReX": "13kDtpZfeE78rUEX50qO-1JYZqYxgJPIKwIyCDkCDH3s",
    "Bandlassian": "1WdRTelFiFVAHDGStc1lxRVyhc5u9dn9ax1Lk1M-akW8",
    "Canvanauts": "17NXjACPkFnfVskxzNZCTRbII2xrDVCOSez4jvBjAYqY",
    "Amakazaam!": "1j1ol8AKbWQOVr2-PfTfjyT2YKO2YBtgOhMsoF72pB9U",
    "V2 Voyagers": "1RByvEr_wjXjokN81PsBnh-N5d51IOKCbXBaMItnDU4E",
}
DRUMS = "1–10"

# Each band: plot items are (kind, label, x, y), the item's centre in metres from house left and
# from the back wall. `backline` labels the riser and three amps (None = not used this set). `mics`
# labels Vox 1–4 (None = fold down); `mic_x` moves them from the even default spacing.
# `patch` is channel(s) -> who / notes. `set` rows are
# (step, song or what happens, seconds, on stage · lead · moves, screen · playback, lighting).
BANDS = [
    {
        "letter": "S",
        "name": "ShipReX",
        "company": "Rex Software",
        "slot": "7:15–7:40 pm",
        "check": "6:00 pm",
        "contact": "Scott Warren",
        "backline": {
            "fender": "Fender · mic ch 13\nScotty",
            "kit": "DW kit · ch 1–10\nSimon",
            "ampeg": None,
            "ac30": "Vox AC30 · mic ch 15\nThiago",
        },
        "mics": ["Vox 1 · ch 21", "Vox 2 · ch 22", "Vox 3 · ch 23", "Vox 4 · ch 24"],
        "mic_x": (2.7, 4.1, 6.3, 7.5),
        "items": [
            ("keys", "Keys · Daniel · TBC\nDI ch 19/20", 1.2, 2.6),
            ("person", "Scotty\nguitar (from\nsong 2)", 1.3, 3.95),
            ("person", "Scott W\nvox", 2.7, SINGER_Y),
            ("person", "Sophie\nerhu", 2.95, 2.85),
            ("di", "Erhu DI · ch 18", 3.4, 2.85),
            ("person", "Kirra\nbass", 6.1, 2.85),
            ("di", "Bass DI · ch 11", 6.55, 2.85),
            ("person", "Hamish\nvox", 4.1, SINGER_Y),
            ("person", "Thiago\nguitar", 5.1, 3.95),
            ("person", "Mindy\nvox", 6.0, SINGER_Y),
            ("person", "Elle\nvox", 6.6, SINGER_Y),
            ("person", "Lil\nvox", 7.5, SINGER_Y),
        ],
        "patch": [
            (DRUMS, "Drum kit", "Simon"),
            ("11", SOURCE[11], "Kirra, DI only (no cab mic)"),
            ("13", SOURCE[13], "Scotty, joins from song 2"),
            ("15", SOURCE[15], "Thiago"),
            ("18", SOURCE[18], "Sophie, erhu pickup"),
            ("19/20", "Keys L/R", "Daniel"),
            ("21", "Vox 1", "Scott W: lead on songs 3 and 4"),
            ("22", "Vox 2", "Hamish: lead on songs 1, 2, 5 and 6"),
            ("23", "Vox 3", "Mindy and Elle share"),
            ("24", "Vox 4", "Lil"),
        ],
        "notes": [
            ("Monitors and in-ears", "No in-ear feeds requested."),
            ("Screen and playback", "None."),
            (
                "Changeover into this set",
                "Opening set, so the stage stays set from their 6:00 pm soundcheck. 11 people on stage: five "
                "singers on four mics, two guitars, bass, drums, keys and erhu.",
            ),
            (
                "Notes and to confirm",
                "Stage layout is as they stood in Brisbane. Is Daniel bringing a keyboard, or using "
                "the Korg Kronos? Lighting for Mustang Sally (new for Sydney); the other lighting notes are carried "
                "over from Brisbane.",
            ),
        ],
        "set": [
            (
                "Song 1",
                "Uprising — Muse (opens with 30 s of Careless Whisper on erhu)",
                360,
                "Everyone except Scotty (guitar). Lead: Hamish (Vox 2)",
                "",
                "Dark, spotlight on the erhu to begin. Blackout on the last erhu note, then hit the full band in "
                "deep red on the Uprising riff. Moody red through the verses with a slow pulse on the stomp beat, "
                "brighter white hits on the choruses",
            ),
            (
                "Song 2",
                "Espresso — Sabrina Carpenter",
                180,
                "Everyone (Scotty joins on guitar). Lead: Hamish (Vox 2)",
                "",
                "Bright and poppy: full warm wash, yellows, oranges and hot pink. No dark moments. Colour sweep or "
                "mirror ball on each chorus; keep the crowd lit",
            ),
            (
                "Song 3",
                "Dumb Things — Paul Kelly",
                180,
                "Everyone. Lead: Scott W (Vox 1)",
                "",
                "Warm white pub-rock wash, even and front-lit so faces are visible. Lift the wash and bring up the "
                "crowd for the singalong choruses",
            ),
            ("Song 4", "Mustang Sally — Wilson Pickett", 250, "Everyone. Lead: Scott W (Vox 1)", "", "Not given yet"),
            (
                "Song 5",
                "Everlong — Foo Fighters",
                270,
                "Everyone. Lead: Hamish (Vox 2)",
                "",
                "Cool blues and purples on the verses, slam to bright white on every chorus. Drop to a single colour "
                "for the quiet bridge, then full blast for the last chorus",
            ),
            (
                "Song 6",
                "Covered in Chrome — Violent Soho",
                225,
                "Everyone. Lead: Hamish (Vox 2)",
                "",
                "Closer, go big: hot white and red, strobes on the chant, everything on for the outro. Hard blackout "
                "on the final hit",
            ),
        ],
        "set_note": "No walk-on or between-song details given.",
    },
    {
        "letter": "B",
        "name": "Bandlassian",
        "company": "Atlassian",
        "slot": "7:50–8:15 pm",
        "check": "5:45 pm",
        "contact": "Steve Kraynov",
        "backline": {
            "fender": "Fender · mic ch 13\nEren / Kazuki",
            "kit": "DW kit · ch 1–10\nEthan",
            "ampeg": "Ampeg\nch 11 + 12",
            "ac30": "Vox AC30 · mic ch 15\nSteve / Josh",
        },
        "mics": [None, "Vox 2 · ch 22", "Vox 3 · ch 23", None],
        "items": [
            ("iem", "Ethan's in-ear amp\nXLR mix feed", 5.2, 1.75),
            ("person", "Eren /\nKazuki\nguitar", 2.1, 2.8),
            ("person", "Linda\nbass", 6.1, 2.85),
            ("person", "Josh /\nMendel\nkeys", 8.2, 2.05),
            ("keys", "Keys (Kronos) · DI ch 19/20", 8.15, 2.6),
            ("person", "Steve /\nJosh\nguitar", 7.3, 3.1),
            ("laptop", "Laptop (Ant) · position TBC\nHDMI video + audio", 0.75, 2.0),
            ("person", "Lesley\nvox", 3.4, SINGER_Y),
            ("person", "Luke\nvox", 5.6, SINGER_Y),
        ],
        "patch": [
            (DRUMS, "Drum kit", "Ethan"),
            ("11", SOURCE[11], "Linda: Ampeg DI out, own Zoom B1 Four pedal in front"),
            ("12", SOURCE[12], "Ampeg cab"),
            ("13", SOURCE[13], "Eren / Kazuki, through their Zoom G5n pedalboard"),
            ("15", SOURCE[15], "Steve / Josh, through their Boss ME-70 pedalboard"),
            ("19/20", "Keys L/R", "Josh / Mendel: Korg Kronos, built-in sounds"),
            ("22", "Vox 2", "Lesley: lead on songs 1 and 2"),
            ("23", "Vox 3", "Luke: lead on songs 4 and 5"),
            ("27/28", "Playback L/R", "Laptop audio over HDMI: walk-on and between songs"),
        ],
        "notes": [
            (
                "Monitors and in-ears",
                "Drummer (Ethan) brings his own in-ear amp and needs one XLR mix feed at the kit.",
            ),
            (
                "Screen and playback",
                "Laptop run by Anthony (“Ant”). Video over HDMI all the way through; audio over HDMI for the "
                "walk-on and the music between songs (ch 27/28).",
            ),
            (
                "Changeover into this set",
                "Strike the erhu DI. Keys go house right beside the AC30. Fold down Vox 1 and Vox 4. Bass moves "
                "from DI to the Ampeg (ch 11 DI + ch 12 mic). Pedalboards: Boss ME-70 at the AC30, Zoom G5n at the "
                "Fender, Zoom B1 Four at the Ampeg. Connect the laptop HDMI and the drummer's XLR feed.",
            ),
            (
                "Notes and to confirm",
                "Where the laptop sits. Six players on stage at a time; guitar and keys players swap between songs "
                "(10–15 s each, see the set). Soundcheck plan: Pokémon intro and solo, Still into You intro and "
                "verse, Life Is a Highway solo, to set the solo boost levels.",
            ),
        ],
        "set": [
            (
                "Walk-on",
                "Walk on, quick band intro",
                20,
                "",
                "Video and audio from HDMI",
                "House lights while the band sets up, then dim for the video",
            ),
            (
                "Song 1",
                "good 4 u — Olivia Rodrigo (rock/metal version)",
                190,
                "Steve (AC30), Eren (Fender), Linda (bass), Josh (keys), Lesley (Vox 2), Luke (Vox 3). Lead: Lesley",
                "Video only",
                "Dim to reds to yellow (standard). Blue for slow bridge?",
            ),
            (
                "1 → 2",
                "",
                15,
                "Eren off, Kazuki on (Fender). Josh off, Mendel on (keys)",
                "Video and changeover music from HDMI",
                "",
            ),
            (
                "Song 2",
                "Still into You — Paramore",
                210,
                "Steve (AC30), Kazuki (Fender), Linda, Mendel (keys), Lesley, Luke. Lead: Lesley",
                "Video only",
                "Blackouts to purples to blue",
            ),
            (
                "2 → 3",
                "",
                15,
                "Kazuki off, Eren on (Fender). Mendel off, Josh on (keys)",
                "Video and audio from HDMI",
                "",
            ),
            (
                "Song 3",
                "Life Is a Highway — Rascal Flatts",
                270,
                "Steve (AC30), Eren (Fender), Linda, Josh (keys), Lesley, Luke. Dual lead vocal. Eren's solo: "
                "boost from his own pedalboard",
                "Video only",
                "Blackouts to yellows. House lights for the second-to-last chorus, when the audience is clapping",
            ),
            (
                "3 → 4",
                "",
                15,
                "Josh moves from keys to guitar on the AC30 (Steve off). Mendel on keys. Eren off, Kazuki on (Fender)",
                "Video and audio from HDMI",
                "",
            ),
            (
                "Song 4",
                "I'm a Believer — Smash Mouth",
                185,
                "Josh (AC30), Kazuki (Fender), Linda, Mendel (keys), Lesley, Luke. Lead: Luke",
                "Video only",
                "Blackouts to greens to yellows to standard",
            ),
            (
                "4 → 5",
                "",
                15,
                "Kazuki off, Eren on (Fender). Kazuki and Steve fetch props and come back on as supporters",
                "Playback audio",
                "",
            ),
            (
                "Song 5",
                "Pokémon Theme — Jason Paige",
                200,
                "Josh (AC30), Eren (Fender), Linda, Mendel (keys), Lesley, Luke. Lead: Luke. Josh's solos: boost "
                "from the ME-70. Pokéball beach balls thrown into the crowd",
                "Video only",
                "Blackouts to blue to yellow (standard), bright yellow snaps during the POKEMON chorus, moving to a "
                "rainbow colour wash for the big finish",
            ),
            ("Finish", "Walk off, pack down", None, "", "", ""),
        ],
        "set_note": "Band expects 20–23 minutes end to end.",
    },
    {
        "letter": "C",
        "name": "Canvanauts",
        "company": "Canva",
        "slot": "8:25–8:50 pm",
        "check": "5:30 pm",
        "contact": "not given yet",
        "backline": {
            "fender": None,
            "kit": "DW kit · ch 1–10\nGerónimo",
            "ampeg": "Ampeg\nJeremy",
            "ac30": None,
        },
        "mics": ["Vox 1 · ch 21\n(over the keys)", "Vox 2 · ch 22", "Vox 3 · ch 23", None],
        "items": [
            ("person", "Danny\nguitar\nTBC", 2.3, 2.75),
            ("di", "Gtr DI · ch 14", 2.75, 2.75),
            ("person", "Ishraque\nguitar\nTBC", 4.5, 2.85),
            ("di", "Gtr DI · ch 17", 4.95, 2.85),
            ("person", "Jeremy\nbass", 6.1, 2.85),
            ("di", "Quad Cortex · ch 11", 6.55, 2.85),
            ("person", "Dave\nguitar\nTBC", 7.85, 2.7),
            ("di", "Gtr DI · ch 16", 7.7, 3.15),
            ("person", "Simone\nkeys + vox", 1.3, 3.3),
            ("keys", "Keys (Kronos) · DI ch 19/20", 1.3, 3.85),
            ("person", "Ollie\nvox", 3.4, SINGER_Y),
            ("person", "Melad\nvox", 5.6, SINGER_Y),
        ],
        "patch": [
            (DRUMS, "Drum kit", "Gerónimo"),
            ("11", SOURCE[11], "Jeremy: Quad Cortex and the Ampeg; routing to confirm"),
            ("12", SOURCE[12], "Only if the Ampeg cab is in use"),
            ("14", SOURCE[14], "Danny: Boss multi-FX"),
            ("16", SOURCE[16], "Dave: UAFX pedal"),
            ("17", "Gtr 3 DI", "Ishraque: Soran Dual Stomp (on the acoustic DI line)"),
            ("19/20", "Keys L/R", "Simone: Korg Kronos"),
            ("21", "Vox 1", "Simone, at the keys: lead on song 3"),
            ("22", "Vox 2", "Ollie: co-lead"),
            ("23", "Vox 3", "Melad: co-lead"),
        ],
        "notes": [
            (
                "Monitors and in-ears",
                "No in-ear feeds requested. All three guitars are DI only, with no amps on stage, so the guitarists "
                "will need guitar in their wedges.",
            ),
            ("Screen and playback", "A visualiser on screen for every song. Source to confirm."),
            (
                "Changeover into this set",
                "Guitar amps aren't used: mute ch 13 and 15. Three DIs for the amp modellers (ch 14, 16, 17). The "
                "Kronos moves from house right to house left, at Vox 1. Vox 1 up; Vox 4 stays down. Bass: Quad "
                "Cortex into ch 11. Strike Bandlassian's laptop.",
            ),
            (
                "Notes and to confirm",
                "Which guitarist stands where. Bass routing (Quad Cortex direct, or through the Ampeg?). Visualiser "
                "source. Contact on the day. The six songs add up to 23:08 before any gaps, in a 25-minute slot.",
            ),
        ],
        "set": [
            (
                "Song 1",
                "The Joker and the Thief — Wolfmother",
                280,
                "Everyone (guitar starts alone, then the rest of the band enters). Lead: Ollie and Melad (Vox 2 + 3)",
                "Visualiser",
                "Red lighting, bright flash when the band kicks in at ~20 s",
            ),
            (
                "Song 2",
                "Thnks fr th Mmrs — Fall Out Boy",
                203,
                "Everyone (guitar and keyboard start with a count-in). Lead: Ollie and Melad",
                "Visualiser",
                "Amber lighting, dramatic, lots of action. Kicks in at ~30 s",
            ),
            (
                "Song 3",
                "So Easy (To Fall in Love) — Olivia Dean",
                170,
                "Everyone (keyboard starts). Lead: Simone (Vox 1)",
                "Visualiser",
                "Purple soft lighting",
            ),
            (
                "Song 4",
                "Beggin' — Måneskin",
                210,
                "Everyone (Vox 2 starts). Lead: Ollie and Melad. Bass solo at ~2:15",
                "Visualiser",
                "Spotlight on vocalist for intro. Red lights on guitar drop after intro piano and vocals. Big stops "
                "at ~0:40, ~1:35, ~2:15 and 2:43",
            ),
            (
                "Song 5",
                "Smooth Criminal — Alien Ant Farm",
                210,
                "Everyone (drums start). Lead: Ollie and Melad",
                "Visualiser",
                "Dynamic white lighting with the stops at ~1:10 and ~2:05",
            ),
            (
                "Song 6",
                "Welcome to the Black Parade — My Chemical Romance",
                315,
                "Everyone (keyboard starts). Lead: Ollie and Melad",
                "Visualiser",
                "Spotlight on vocalist for intro. Red lights on guitar drop after intro piano and vocals",
            ),
        ],
        "set_note": "No walk-on or between-song details given.",
    },
    {
        "letter": "A",
        "name": "Amakazaam!",
        "company": "Amazon",
        "slot": "9:00–9:25 pm",
        "check": "5:15 pm",
        "contact": "Keira Daley",
        "backline": {
            "fender": "Fender · mic ch 13\nTim",
            "kit": "DW kit · ch 1–10\nAndrew (Keira on song 4)",
            "ampeg": "Ampeg\nch 11 + 12",
            "ac30": "Vox AC30 · mic ch 15\nTyler",
        },
        "mics": ["Vox 1 · ch 21\n(over the keys)", "Wireless · ch 26\nstraight stand", "Vox 3 · ch 23", None],
        "items": [
            ("mic", "Vox 5 · ch 25", 3.95, 1.75),
            ("person", "Tim\nrhythm\nguitar", 2.3, 2.75),
            ("laptop", "Laptop · side stage, TBC\nHDMI video + audio", 0.45, 2.0),
            ("person", "Sam\nbass", 6.1, 2.85),
            ("iem", "Keira's wireless · TBC\nreceiver out · in-ear mix in", 8.4, 2.6),
            ("person", "Aaron\nkeys + vox", 1.3, 3.3),
            ("keys_own", "Keys (own) · stereo DI ch 19/20", 1.3, 3.85),
            ("person", "Keira\nlead vox", 3.4, SINGER_Y),
            ("person", "Tyler\nguitar\n+ vox", 5.6, SINGER_Y),
        ],
        "patch": [
            (DRUMS, "Drum kit", "Andrew; Keira plays on song 4"),
            ("11", SOURCE[11], "Sam: Ampeg DI out (uses a tuner)"),
            ("12", SOURCE[12], "Ampeg cab"),
            ("13", SOURCE[13], "Tim, with pedalboard"),
            ("15", SOURCE[15], "Tyler, with pedalboard"),
            ("19/20", "Keys L/R", "Aaron: own keyboard, stereo, needs a DI with two inputs"),
            ("21", "Vox 1", "Aaron: backing on song 6. Andrew's cowbell spot on song 4"),
            ("23", "Vox 3", "Tyler: co-lead on song 3, backing on song 6"),
            ("25", "Vox 5 (drums)", "Andrew: spoken line. Keira: sings song 4 from the kit"),
            ("26", SOURCE[26], "Keira: own mic and receiver (XLR out), on a straight stand at centre"),
            ("27/28", "Playback L/R", "Laptop audio over HDMI: Alexa voice-over"),
        ],
        "notes": [
            (
                "Monitors and in-ears",
                "Keira has her own wireless in-ears; her transmitter takes an XLR from the desk. The mix is vocals "
                "only: her wireless (ch 26), plus Vox 5 when she sings from the kit (song 4), plus a little Vox 3 "
                "on song 3.",
            ),
            (
                "Screen and playback",
                "Laptop at side stage on HDMI. Video all the way through; Alexa voice-over audio at the intro, "
                "between songs 3 → 4 and 5 → 6, and at the finish (ch 27/28). Keira runs it with a clicker.",
            ),
            (
                "Changeover into this set",
                "Strike the Kronos and the three guitar DIs. Aaron's keyboard and stand go house left at Vox 1, on "
                "a stereo DI. Amps back in: Tim on the Fender, Tyler on the AC30. Swap the Vox 2 stand for a "
                "straight (boomless) stand with Keira's wireless mic and clip. Vox 5 up at the kit. Patch her "
                "receiver (ch 26) and in-ear transmitter. Andrew fits clamps for cowbell and tambourine.",
            ),
            (
                "Notes and to confirm",
                "Song 4: Keira and Andrew swap on the kit. Andrew plays cowbell at Vox 1, then takes the kit back "
                "at the guitar solo. Where the wireless receiver and in-ear transmitter sit. Planned running time "
                "is 24:00 in a 25-minute slot.",
            ),
        ],
        "set": [
            (
                "Intro",
                "Alexa's intro (band already on stage)",
                10,
                "",
                "Laptop audio and video",
                "Blackout until Alexa says “Amakazaam!”, then bring up a warm wash",
            ),
            (
                "Song 1",
                "Valerie — Amy Winehouse",
                210,
                "Everyone. Lead: Keira (wireless, centre)",
                "Laptop video",
                "Warm vintage vibes, gold and pink",
            ),
            ("1 → 2", "Straight into the next song", 30, "Keira talks over the intro", "Laptop video", "Wash with screen visible"),
            (
                "Song 2",
                "Superstition — Stevie Wonder",
                180,
                "Everyone. Lead: Keira",
                "Laptop video",
                "Groovy and spooky? Halloween orange, occasional greens, occasional disco?",
            ),
            ("2 → 3", "Straight into the next song", 20, "Keira talks over the intro", "Laptop video", "Wash with screen visible"),
            (
                "Song 3",
                "Under Pressure — Queen & David Bowie",
                210,
                "Everyone. Lead: Keira (wireless) and Tyler (Vox 3). Add a bit of Vox 3 to Keira's in-ears",
                "Laptop video",
                "Start with spotlights on lead vox (wireless) and lead guitar (Vox 3). Bright whites with red and "
                "blue accents. Build and break tension",
            ),
            (
                "3 → 4",
                "Straight into the next song, over dialogue between Keira, Alexa and Andrew",
                30,
                "Andrew speaks from the kit (Vox 5), then Keira takes over the drums and Andrew moves to cowbell at "
                "Vox 1. Keira's in-ears need Vox 5",
                "Laptop audio and video",
                "Wash with screen visible",
            ),
            (
                "Song 4",
                "Are You Gonna Go My Way — Lenny Kravitz",
                180,
                "Everyone. Keira sings from the kit on Vox 5; Andrew on cowbell near Vox 1. At the hits at the top "
                "of the guitar solo Andrew takes the kit back, and Keira sings the final line on the wireless at "
                "centre",
                "Laptop video",
                "Cue: Keira's first drum fill (after the opening hits). Epic stadium: flashes on song accents, "
                "backlights, colour and movement",
            ),
            (
                "4 → 5",
                "Straight into the next song",
                10,
                "Keira talks over the intro. Sam tunes down",
                "Laptop video",
                "Moody, but with screen visible",
            ),
            (
                "Song 5",
                "You Oughta Know — Alanis Morissette",
                240,
                "Everyone. Lead: Keira",
                "Laptop video",
                "At the first lyric, bring the wash down to just a spotlight on lead vox. Moody blue and magenta. "
                "Start small and build to hectic",
            ),
            (
                "5 → 6",
                "Alexa dialogue",
                10,
                "Keira (wireless) and Alexa (laptop). Sam tunes back up",
                "Laptop audio and video",
                "Wash with screen visible",
            ),
            (
                "Song 6",
                "Killing in the Name — Rage Against the Machine",
                300,
                "Everyone. Lead: Keira. Backing: Tyler (Vox 3), Aaron (Vox 1)",
                "Laptop video: the on-screen gauge explodes at the big section",
                "Song cue: “Okay crew, we know what to do.” Start dark, end in absolute carnage: red and orange "
                "like flames, flashes on song accents, strobe by the end. Important: “destruction” lights at the "
                "big “fuck you I won't do what you tell me!” section",
            ),
            ("Finish", "Alexa's outro, curtain call", 10, "", "Laptop audio and video", "Wash with screen visible"),
        ],
        "set_note": "The band's own run sheet has the full spoken script.",
    },
    {
        "letter": "V",
        "name": "V2 Voyagers",
        "company": "V2 AI",
        "slot": "9:35–10:00 pm",
        "check": "5:00 pm",
        "contact": "Holly or Shane",
        "backline": {
            "fender": "Fender · mic ch 13\nTaz",
            "kit": "DW kit · ch 1–10\nJai (songs 1–3) · Shane (4–5)",
            "ampeg": "Ampeg\nch 11 + 12",
            "ac30": None,
        },
        "mics": [None, "Vox 2 · ch 22\nHolly's own Beta 58", None, "Vox 4 · ch 24"],
        "items": [
            ("laptop", "Laptop on the riser\nHDMI video · click to drummer", 3.75, 1.65),
            ("di", "Backing DI · ch 27", 4.2, 2.55),
            ("person", "Taz\nguitar", 2.2, 2.75),
            ("person", "Cam\nbass", 6.1, 2.85),
            ("iem", "Holly's in-ears (Xvive U4)\ntransmitter needs a mix feed", 2.3, 3.55),
            ("di", "Acoustic DI · ch 17", 7.05, 3.3),
            ("person", "Holly\nlead vox", 3.4, SINGER_Y),
            ("person", "Milly\nacoustic\n+ vox", 7.7, SINGER_Y),
        ],
        "patch": [
            (DRUMS, "Drum kit", "Jai on songs 1–3 (lighter player), Shane on songs 4–5 (hard hitter)"),
            ("11", SOURCE[11], "Cam: Ampeg DI out"),
            ("12", SOURCE[12], "Ampeg cab"),
            ("13", SOURCE[13], "Taz"),
            ("17", SOURCE[17], "Milly"),
            ("22", "Vox 2", "Holly, on her own Shure Beta 58"),
            ("24", "Vox 4", "Milly: light backing vocals, and tambourine"),
            ("27", "Playback L (mono)", "Backing track from the band's Mooer Micro DI at the riser"),
        ],
        "notes": [
            (
                "Monitors and in-ears",
                "Holly has her own wireless in-ears (Xvive U4); the transmitter needs a mix feed from the desk. The "
                "drummers' click comes from the band's own headphone amp, not the desk.",
            ),
            (
                "Screen and playback",
                "The band's laptop at the drum riser runs the whole set as one Reaper session. Video over HDMI (no "
                "audio), already cut and cued, so no VJ needed. Backing track is mono, from their own Mooer Micro DI "
                "(or a house DI) into ch 27. An idle loop plays between songs.",
            ),
            (
                "Changeover into this set",
                "Strike Aaron's keys, Keira's wireless rig and the side-stage laptop. Regular stand back at Vox 2, "
                "with Holly's own Beta 58. AC30 isn't used: mute ch 15. Acoustic DI (ch 17) at Vox 4. Fold down "
                "Vox 1, 3 and 5. HDMI run and a DI line to the laptop at the riser.",
            ),
            (
                "Notes and to confirm",
                "Two drummers; about a minute to swap before song 4. Milly chose Vox 4 (house right) but also wrote "
                "“stage right”, which is house left: confirm her side. The band asks for a smoke machine on "
                "the walk-on.",
            ),
        ],
        "set": [
            (
                "Walk-on",
                "Intro video and audio",
                30,
                "",
                "Intro audio and video",
                "No front lighting, band in silhouette. Bring lights up as the rocket takes off. Smoke machine "
                "here too",
            ),
            (
                "Song 1",
                "Venus — Bananarama",
                160,
                "Holly, Cam, Taz, Milly, Jai. Lead: Holly",
                "Backing audio and video",
                "Match colours of video: reds, blues, purples",
            ),
            ("1 → 2", "", 15, "", "Idle loop video", ""),
            (
                "Song 2",
                "Cosmic Girl — Jamiroquai",
                220,
                "Holly, Cam, Taz, Milly, Jai. Lead: Holly",
                "Backing audio and video",
                "Match colours of video: amber, sky blue, purples",
            ),
            ("2 → 3", "", 15, "", "Idle loop video", ""),
            (
                "Song 3",
                "Drops of Jupiter — Train",
                220,
                "Holly, Cam, Taz, Milly, Jai. Lead: Holly",
                "Backing audio and video",
                "Match colours of video: turquoise, deep blue, lime green",
            ),
            ("3 → 4", "Drummer changeover", 60, "Jai off, Shane on", "Idle loop video", ""),
            (
                "Song 4",
                "All Star — Smash Mouth",
                200,
                "Holly, Cam, Taz, Milly, Shane. Lead: Holly",
                "Backing audio and video",
                "Match colours of video: deep greens, reds, sky blue",
            ),
            ("4 → 5", "", 15, "", "Idle loop video", ""),
            (
                "Song 5",
                "Medley: Starships — Nicki Minaj / UFO — Sneaky Sound System",
                320,
                "Holly, Cam, Taz, Milly, Shane. Lead: Holly",
                "Backing audio and video",
                "Starships: purples, pinks, blues. UFO: golds, reds, whites",
            ),
        ],
        "set_note": "",
    },
]


def rgb(hex_):
    return RGBColor.from_string(hex_)


def write(tf, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT, lead=None):
    """Fill a text frame; `lead` is a bold run in front of the first paragraph's text."""
    tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        parts = [(lead, True)] if (lead and i == 0) else []
        parts.append((line, bold))
        for chunk, b in parts:
            r = p.add_run()
            r.text = chunk
            r.font.size, r.font.bold, r.font.name = Pt(size), b, FONT
            r.font.color.rgb = rgb(color)


def text(slide, x, y, w, h, body, size=9, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    write(tf, body, size, color, bold, align)
    return tb


def shape(slide, x, y, w, h, body, fill, line, size=7, color=INK, dashed=False, kind=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(line)
    sh.line.width = Pt(1)
    if dashed:
        sh.line.dash_style = MSO_LINE.DASH
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    write(tf, body, size, color, align=PP_ALIGN.CENTER)
    sh.name = body.split("\n")[0]
    return sh


def lines_needed(body, width, size):
    """Rough wrapped-line count for Arial at `size` pt in a `width`-inch column."""
    per_line = max(1, int((width - 0.12) * 72 / (size * 0.5)))
    return sum(max(1, ceil(len(part) / per_line)) for part in str(body).split("\n"))


def row_height(cells, widths, size):
    return max(lines_needed(c, w, size) for c, w in zip(cells, widths)) * size * 1.22 / 72 + 0.11


def border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = OxmlElement(tag)
        ln.set("w", "6350")
        fill = OxmlElement("a:solidFill")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", RULE)
        fill.append(clr)
        ln.append(fill)
        tc_pr.append(ln)


def table(slide, x, y, widths, header, rows, size=8, fills=None, bold_cols=(0,), center_cols=()):
    """Draw a table and return the y of its bottom edge. `fills` maps row index -> fill hex."""
    heights = [row_height(header, widths, size)] + [row_height(r, widths, size) for r in rows]
    frame = slide.shapes.add_table(
        len(rows) + 1, len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(sum(heights))
    )
    tbl = frame.table
    tbl.horz_banding = False
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for r, cells in enumerate([header] + list(rows)):
        tbl.rows[r].height = Inches(heights[r])
        for c, val in enumerate(cells):
            cell = tbl.cell(r, c)
            border(cell)  # borders first: the fill element must follow them in the cell XML
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(PALE if r == 0 else (fills or {}).get(r - 1, "FFFFFF"))
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.035)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE if r == 0 else MSO_ANCHOR.TOP
            write(
                cell.text_frame,
                val,
                size,
                bold=r == 0 or c in bold_cols,
                align=PP_ALIGN.CENTER if c in center_cols else PP_ALIGN.LEFT,
            )
    return y + sum(heights)


def page(prs, title, right=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text(slide, MARGIN, 0.22, 7.2, 0.4, title, size=17, bold=True)
    if right:
        text(slide, 6.2, 0.2, PAGE_W - MARGIN - 6.2, 0.42, right, size=9, color=GREY, align=PP_ALIGN.RIGHT)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(0.68), Inches(PAGE_W - 2 * MARGIN), Inches(0.03))
    bar.shadow.inherit = False
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(BRAND)
    bar.line.fill.background()
    return slide


def clock(seconds):
    return f"{seconds // 60}:{seconds % 60:02d}"


def running_time(band):
    known = sum(s for _, _, s, *_ in band["set"] if s)
    missing = [what.split(" — ")[0] for _, what, s, *_ in band["set"] if s is None and what and "Walk off" not in what]
    total = f"{clock(known)} planned"
    if missing:
        total += f", plus {' and '.join(missing)} (length not given)"
    return total


def overview(prs):
    slide = page(prs, "Stage plots, patches and sets", EVENT)
    text(slide, MARGIN, 0.8, 5.6, 0.3, "Schedule", size=11, bold=True, color=BRAND)
    fills = {i: BETWEEN for i, (_, what, _) in enumerate(SCHEDULE) if what == "Changeover"}
    table(slide, MARGIN, 1.1, [1.0, 2.35, 2.3], ["Time", "What", "Notes"], SCHEDULE, size=9, fills=fills)
    x, y = 6.45, 0.8
    text(slide, x, y, 4.9, 0.3, "How to read these", size=11, bold=True, color=BRAND)
    y += 0.34
    for head, body in HOW_TO_READ:
        h = (lines_needed(head + ". " + body, 4.9, 9.5) + 0.6) * 9.5 * 1.22 / 72
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(4.9), Inches(h))
        tb.text_frame.margin_left = tb.text_frame.margin_top = Inches(0.03)
        write(tb.text_frame, body, 9.5, lead=head + ". ")
        y += h + 0.1
    y += 0.1
    text(slide, x, y, 4.9, 0.3, "Band stage plots and sets", size=11, bold=True, color=BRAND)
    y += 0.34
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(4.9), Inches(1.6))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = Inches(0.03)
    write(tf, "One deck per band, in show order. Drafts until each band confirms.", 9.5)
    for band in BANDS:
        para = tf.add_paragraph()
        para.space_before = Pt(3)
        lead = para.add_run()
        lead.text = f"{band['letter']}  {band['slot'].replace(' pm', '')}  "
        lead.font.size, lead.font.name, lead.font.bold = Pt(9.5), FONT, True
        lead.font.color.rgb = rgb(INK)
        link = para.add_run()
        link.text = f"{band['name']} ({band['company']})"
        link.font.size, link.font.name, link.font.underline = Pt(9.5), FONT, True
        link.font.color.rgb = rgb(BRAND)
        link.hyperlink.address = DRIVE.format(DECKS[band["name"]])


def changeovers(prs):
    slide = page(prs, "Changeovers", "What changes on stage going into each set")
    rows = [
        (band["slot"], f"{band['name']} ({band['company']})", body)
        for band in BANDS
        for title, body in band["notes"]
        if title.startswith("Changeover")
    ]
    table(slide, MARGIN, 0.9, [1.2, 2.0, PAGE_W - 2 * MARGIN - 3.2], ["Set", "Into", "What changes"], rows, size=9.5)


def festival_patch(prs):
    slide = page(prs, "Festival patch", "The same 32 inputs for every band · ● used · ○ only if needed")
    letters = [b["letter"] for b in BANDS]
    used = {}
    for band in BANDS:
        for chans, _, who in band["patch"]:
            lo, _, hi = chans.replace("/", "–").partition("–")
            for ch in range(int(lo), int(hi or lo) + 1):
                used[(band["letter"], ch)] = "○" if who.startswith("Only if") else "●"
    widths = [0.36, 1.45, *[0.27] * 5, 2.05]
    for col, chunk in enumerate((FESTIVAL[:16], FESTIVAL[16:])):
        rows = [(str(ch), src, *[used.get((ltr, ch), "") for ltr in letters], note) for ch, src, note in chunk]
        table(
            slide,
            MARGIN + col * (sum(widths) + 0.27),
            0.9,
            widths,
            ["Ch", "Source", *letters, "Notes"],
            rows,
            size=8,
            center_cols=range(2, 7),
        )
    text(
        slide,
        MARGIN,
        7.45,
        PAGE_W - 2 * MARGIN,
        0.5,
        "S ShipReX · B Bandlassian · C Canvanauts · A Amakazaam! · V V2 Voyagers. Walk-in and changeover music "
        "on the desk's stereo line inputs, outside the 32. The host uses one of the band vocal mics between sets.",
        size=9,
        color=GREY,
    )


def item(slide, kind, label, x, y, scale, x0, y0, w, used=True):
    """One plot item, grouped with its label so it drags as one piece. (x, y) is its centre in metres."""
    fill, line, wm, dm, where = STYLE[kind]
    sw, sh = wm * scale, dm * scale
    left, top = x0 + x * scale - sw / 2, y0 + y * scale - sh / 2
    if not used:
        fill, line = "FFFFFF", RULE
    grp = slide.shapes.add_group_shape()
    grp.name = label.split("\n")[0]
    body = grp.shapes.add_shape(
        MSO_SHAPE.OVAL if kind in ("person", "mic") else MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(sw),
        Inches(sh),
    )
    body.shadow.inherit = False
    body.fill.solid()
    body.fill.fore_color.rgb = rgb(fill)
    body.line.color.rgb = rgb(line)
    body.line.width = Pt(1)
    if not used:
        body.line.dash_style = MSO_LINE.DASH
    n = label.count("\n") + 1
    th = n * 7 * 1.2 / 72 + 0.04
    if where == "inside":
        tw = max(sw, 1.0 if kind == "person" else 1.6)
        box = (left + sw / 2 - tw / 2, top + sh / 2 - th / 2, tw, th, PP_ALIGN.CENTER)
    elif where == "below":
        tw = 1.9
        box = (left + sw / 2 - tw / 2, top + sh + 0.01, tw, th, PP_ALIGN.CENTER)
    else:
        tw = 1.4
        box = (left + sw + 0.04, top + sh / 2 - th / 2, tw, th, PP_ALIGN.LEFT)
    bx, by, bw, bh, align = box
    bx = min(max(bx, x0 + 0.02), x0 + w - bw - 0.02)
    tb = grp.shapes.add_textbox(Inches(bx), Inches(by), Inches(bw), Inches(bh))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    write(tf, label, 7, color=INK if used else "9CA3AF", align=align)


def stage_plot(slide, band, x0, y0, w):
    """Draw the stage to scale at width `w` inches; return its height in inches."""
    scale = w / STAGE_M[0]
    h = STAGE_M[1] * scale
    text(slide, x0, y0 - 0.21, w, 0.2, "UPSTAGE (back wall)", size=7, color=GREY, align=PP_ALIGN.CENTER)
    stage = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0), Inches(y0), Inches(w), Inches(h))
    stage.shadow.inherit = False
    stage.fill.solid()
    stage.fill.fore_color.rgb = rgb("FAFAFA")
    stage.line.color.rgb = rgb("9CA3AF")
    stage.name = "Stage"
    # Below the front mics' labels, which hang just past the stage edge.
    text(slide, x0, y0 + h + 0.24, w, 0.2, "AUDIENCE", size=7, color=GREY, bold=True, align=PP_ALIGN.CENTER)
    text(slide, x0, y0 + h + 0.24, 2.2, 0.2, "◀ house left (stage right)", size=7, color=GREY)
    text(slide, x0 + w - 2.2, y0 + h + 0.24, 2.2, 0.2, "house right (stage left) ▶", size=7, color=GREY, align=PP_ALIGN.RIGHT)

    def place(kind, label, x, y, used=True):
        item(slide, kind, label, x, y, scale, x0, y0, w, used)

    # Drum riser centred against the back wall.
    place("riser", band["backline"]["kit"] + "\nriser 2.4 × 1.8 m", STAGE_M[0] / 2, RISER_Y)
    unused = {"fender": "Fender · not used", "ampeg": "Ampeg · not mic'd", "ac30": "Vox AC30 · not used"}
    for key, (x, y) in BACKLINE.items():
        label = band["backline"][key]
        place(key, label or unused[key], x, y, used=bool(label))
    for kind, label, x, y in band["items"]:
        place(kind, label, x, y)
    for n, (label, x) in enumerate(zip(band["mics"], band.get("mic_x", MIC_X)), start=1):
        place("mic", label or f"Vox {n} · {FOLD}", x, MIC_Y, used=bool(label))
    return h


def band_pages(prs, band, venue=True):
    """Add a band's stage plot and set pages. Band copies leave out the changeover notes, which
    describe the previous band's gear."""
    head = f"{band['name']} ({band['company']})"
    right = f"Set {band['slot']} · soundcheck {band['check']} · contact on the day: {band['contact']}"
    slide = page(prs, f"{head} — stage plot", right)

    plot_w, top = PAGE_W - 2 * MARGIN, 0.85
    plot_bottom = top + 0.22 + stage_plot(slide, band, MARGIN, top + 0.22, plot_w) + 0.46
    lx = MARGIN
    for kind, label in LEGEND:
        fill, line, *_ = STYLE[kind]
        shape(slide, lx, plot_bottom + 0.07, 0.16, 0.13, "", fill, line)
        text(slide, lx + 0.18, plot_bottom + 0.03, 1.0, 0.2, label, size=7, color=GREY)
        lx += 0.3 + len(label) * 0.062
    text(slide, lx, plot_bottom + 0.03, 2.0, 0.2, "Dashed = not used this set", size=7, color=GREY)
    metre = plot_w / STAGE_M[0]
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGIN), Inches(plot_bottom + 0.38), Inches(metre), Inches(0.05))
    bar.shadow.inherit = False
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(INK)
    bar.line.fill.background()
    text(
        slide,
        MARGIN + metre + 0.08,
        plot_bottom + 0.31,
        6.0,
        0.2,
        "1 m · all to scale; DI boxes and mic stands slightly enlarged",
        size=7,
        color=GREY,
    )

    slide = page(prs, f"{head} — patch and notes", right)
    widths = [0.55, 1.25, 3.3]
    bottom = table(slide, MARGIN, top, widths, ["Ch", "Input", "Who and notes"], band["patch"], size=8.5)
    live = set()
    for chans, *_ in band["patch"]:
        lo, _, hi = chans.replace("/", "–").partition("–")
        live.update(range(int(lo), int(hi or lo) + 1))
    idle = [ch for ch in range(11, 29) if ch not in live]
    text(
        slide,
        MARGIN,
        bottom + 0.05,
        sum(widths),
        0.5,
        "Not used this set: ch " + ", ".join(map(str, idle)) + ". Ch 31/32 are the crowd mics, for the recording only.",
        size=8,
        color=GREY,
    )

    notes = [n for n in band["notes"] if venue or not n[0].startswith("Changeover")]
    nx = MARGIN + sum(widths) + 0.3
    nw, ny = PAGE_W - MARGIN - nx, top
    for title, body in notes:
        nh = (lines_needed(body, nw - 0.16, 9.5) * 9.5 + 10) * 1.25 / 72 + 0.2
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(nx), Inches(ny), Inches(nw), Inches(nh))
        box.shadow.inherit = False
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(PALE)
        box.line.fill.background()
        box.name = title
        tf = box.text_frame
        tf.margin_left = tf.margin_right = Inches(0.08)
        tf.margin_top = tf.margin_bottom = Inches(0.06)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        write(tf, title, 10, color=BRAND, bold=True)
        para = tf.add_paragraph()
        r = para.add_run()
        r.text = body
        r.font.size, r.font.name = Pt(9.5), FONT
        r.font.color.rgb = rgb(INK)
        ny += nh + 0.12

    widths = [0.65, 2.35, 0.55, 3.0, 1.55, PAGE_W - 2 * MARGIN - 8.1]
    header = ["Step", "Song, or what happens", "Length", "On stage · lead · moves", "Screen · playback", "Lighting"]
    rows = [(step, what, clock(s) if s else "", stage, screen, light) for step, what, s, stage, screen, light in band["set"]]
    summary = f"Running time: {running_time(band)}, in a 25-minute slot. {band['set_note']}".strip()
    limit, start, part = PAGE_H - 0.45, 0, 1
    while start < len(rows):
        y, end = 1.2 + row_height(header, widths, 8), start
        while end < len(rows) and y + row_height(rows[end], widths, 8) <= limit:
            y += row_height(rows[end], widths, 8)
            end += 1
        slide = page(prs, f"{head} — the set" + (f" ({part})" if part > 1 else ""), right)
        text(slide, MARGIN, 0.8, PAGE_W - 2 * MARGIN, 0.3, summary, size=9, color=GREY)
        fills = {i: BETWEEN for i, r in enumerate(rows[start:end]) if "→" in r[0]}
        table(slide, MARGIN, 1.2, widths, header, rows[start:end], size=8, fills=fills)
        start, part = end, part + 1


def deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(PAGE_W), Inches(PAGE_H)
    return prs


def main():
    OUT_DIR.mkdir(exist_ok=True)
    venue, collated = deck(), deck()
    for prs in (venue, collated):
        overview(prs)
        festival_patch(prs)
        changeovers(prs)
    for n, band in enumerate(BANDS, start=1):
        band_pages(collated, band)
        own = deck()
        band_pages(own, band, venue=False)
        slug = band["name"].replace("!", "").replace(" ", "-")
        own.save(OUT_DIR / f"BOTTB-Sydney-2026-{n}-{slug}-stage-plot-and-set.pptx")
        print(f"  {band['name']}: {running_time(band)} ({len(own.slides)} pages)")
    for prs, name in ((venue, "venue-overview"), (collated, "all-bands-collated")):
        out = OUT_DIR / f"BOTTB-Sydney-2026-{name}.pptx"
        prs.save(out)
        print(out, f"({len(prs.slides)} pages)")


if __name__ == "__main__":
    main()
